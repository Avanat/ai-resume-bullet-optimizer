"""Build downloadable versions of the optimized resume (txt / markdown / docx / json)."""
from __future__ import annotations

import io
import json
from typing import Dict, List, Tuple

from docx import Document
from docx.shared import Inches, Pt

from .resume_parser import EMAIL_RE, LINK_RE, SECTION_TITLES, Bullet, ParsedResume

Line = Tuple[str, str]  # ("bullet" | "text", content)
BOLD_SECTIONS = {"experience", "projects", "education"}


def _contact_line(p: ParsedResume) -> str:
    return " | ".join(x for x in [p.email, p.phone, *p.links] if x)


def build_sections(p: ParsedResume, replacements: Dict[int, str],
                   manual: List[Bullet]) -> List[Tuple[str, str, List[Line]]]:
    """Return [(section_key, title, lines)] with accepted rewrites swapped in."""
    out: List[Tuple[str, str, List[Line]]] = []
    for key, blocks in p.blocks.items():
        lines: List[Line] = []
        for b in blocks:
            if key == "header":
                s = b.text
                if s == p.name or EMAIL_RE.search(s) or LINK_RE.search(s) or (p.phone and p.phone in s):
                    continue
            text = replacements.get(b.bullet_id, b.text) if b.bullet_id is not None else b.text
            lines.append((b.kind, text))
        if lines:
            title = SECTION_TITLES.get(key, key.title())
            out.append((key, "" if key == "header" else title, lines))
    if manual:
        out.append(("manual", "Additional Achievements",
                    [("bullet", replacements.get(m.id, m.text)) for m in manual]))
    return out


def to_text(p: ParsedResume, replacements: Dict[int, str], manual: List[Bullet]) -> str:
    rows = [p.name.upper(), _contact_line(p), ""]
    for _, title, lines in build_sections(p, replacements, manual):
        if title:
            rows += [title.upper(), "-" * len(title)]
        rows += [("• " + t) if k == "bullet" else t for k, t in lines]
        rows.append("")
    return "\n".join(rows).strip() + "\n"


def to_markdown(p: ParsedResume, replacements: Dict[int, str], manual: List[Bullet]) -> str:
    rows = [f"# {p.name}", _contact_line(p), ""]
    for key, title, lines in build_sections(p, replacements, manual):
        if title:
            rows += [f"## {title}", ""]
        for k, t in lines:
            if k == "bullet":
                rows.append(f"- {t}")
            else:
                rows += ["", f"**{t}**" if key in BOLD_SECTIONS else t, ""]
        rows.append("")
    out = "\n".join(rows)
    while "\n\n\n" in out:
        out = out.replace("\n\n\n", "\n\n")
    return out.strip() + "\n"


def to_docx(p: ParsedResume, replacements: Dict[int, str], manual: List[Bullet]) -> bytes:
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.8)
        s.top_margin = s.bottom_margin = Inches(0.7)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    doc.add_heading(p.name or "Resume", level=1)
    if _contact_line(p):
        doc.add_paragraph(_contact_line(p))
    for key, title, lines in build_sections(p, replacements, manual):
        if title:
            doc.add_heading(title, level=2)
        for kind, text in lines:
            if kind == "bullet":
                doc.add_paragraph(text, style="List Bullet")
            else:
                para = doc.add_paragraph()
                run = para.add_run(text)
                run.bold = key in BOLD_SECTIONS
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def to_json(payload: dict) -> str:
    return json.dumps(payload, indent=2, default=str)
